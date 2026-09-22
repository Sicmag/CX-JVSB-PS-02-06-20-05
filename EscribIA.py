"""
EscribIA - Prototipo funcional (versión terminal)
Convierte apuntes manuscritos en documentos usando IA.

Salidas disponibles:
  1. Documento de Word (.docx)
  2. Presentación de PowerPoint (.pptx) con explicaciones
  3. Resumen corto de estudio (.docx)

Proyecto formativo SENA.
"""

import base64
import os
import time
from datetime import datetime

import requests
from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

from config import GEMINI_API_KEY, GROQ_API_KEY


# ============ CONFIGURACIÓN ============

MODELO_GEMINI = "gemini-3.6-flash"
URL_GEMINI = (
    f"https://generativelanguage.googleapis.com/v1beta/models/"
    f"{MODELO_GEMINI}:generateContent?key={GEMINI_API_KEY}"
)

MODELO_GROQ = "qwen/qwen3.8-27b"
URL_GROQ = "https://api.groq.com/openai/v1/chat/completions"


PROMPT_WORD = """Eres un asistente que digitaliza apuntes manuscritos en español.

Recibirás la foto de un apunte escrito a mano. Tu tarea:

1. Transcribe el texto con la mayor fidelidad posible.
2. Corrige errores ortográficos evidentes sin cambiar el sentido.
3. Organiza el contenido en secciones con títulos claros.
4. Usa listas con guiones cuando el apunte tenga enumeraciones.
5. Si hay fórmulas, dibujos o partes ilegibles, indícalo entre corchetes.
6. Devuelve SOLO el texto final, sin comentarios, sin markdown con asteriscos.
"""


PROMPT_PPT = """Eres un asistente que convierte apuntes manuscritos en presentaciones.

Recibirás la foto de un apunte escrito a mano. Organízalo en diapositivas y
para cada una agrega una breve explicación del tema.

Devuelve el contenido en ESTE FORMATO EXACTO, sin nada más:

TITULO: <título corto, máximo 8 palabras>
EXPLICACION: <una o dos frases explicando el tema, en tus propias palabras, sin inventar información que no esté en el apunte>
- <punto clave 1>
- <punto clave 2>
- <punto clave 3>
---
TITULO: <título de la siguiente diapositiva>
EXPLICACION: <explicación del tema>
- <punto 1>
- <punto 2>
---

Reglas:
- Entre 4 y 8 diapositivas máximo.
- Cada diapositiva con 3 a 5 puntos.
- Cada punto de máximo 12 palabras.
- La explicación debe ser breve (1-2 líneas) y en tono didáctico.
- No uses asteriscos ni markdown.
- Empieza directo con "TITULO:".
"""


PROMPT_RESUMEN = """Eres un asistente que crea resúmenes de estudio.

Recibirás la foto de un apunte manuscrito en español. Tu tarea:

1. Crea un resumen corto y claro (máximo 300 palabras).
2. Empieza con un párrafo de 2-3 líneas explicando de qué trata.
3. Luego lista los conceptos clave como viñetas.
4. Termina con 3 preguntas de repaso sobre el tema.
5. No uses asteriscos ni markdown especial, solo texto y guiones.
"""


# ============ BÚSQUEDA DE IMAGEN ============

def encontrar_ultima_foto():
    """Busca la foto más reciente en las carpetas típicas del dispositivo."""
    carpetas = [
        "/storage/emulated/0/DCIM/Camera",
        "/storage/emulated/0/Pictures",
        "/storage/emulated/0/Pictures/Screenshots",
        "/storage/emulated/0/Download",
        "/sdcard/DCIM/Camera",
        "/sdcard/Pictures",
    ]
    extensiones = (".jpg", ".jpeg", ".png", ".webp")
    mejor = None
    mejor_tiempo = 0

    for carpeta in carpetas:
        if not os.path.isdir(carpeta):
            continue
        for archivo in os.listdir(carpeta):
            if not archivo.lower().endswith(extensiones):
                continue
            ruta = os.path.join(carpeta, archivo)
            try:
                t = os.path.getmtime(ruta)
                if t > mejor_tiempo:
                    mejor_tiempo = t
                    mejor = ruta
            except OSError:
                continue

    return mejor


# ============ PROCESAMIENTO DE IMAGEN ============

def cargar_imagen(ruta):
    """Convierte la imagen a base64 en JPEG."""
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"No se encontró la imagen: {ruta}")

    img = Image.open(ruta)
    img.thumbnail((1600, 1600))

    temp = "_temp_escribia.jpg"
    img.convert("RGB").save(temp, "JPEG", quality=85)

    with open(temp, "rb") as f:
        data = base64.b64encode(f.read()).decode("utf-8")
    os.remove(temp)

    return data, "image/jpeg"


# ============ PROVEEDOR 1: GEMINI ============

def transcribir_con_gemini(imagen_b64, mime_type, prompt):
    payload = {
        "contents": [
            {
                "parts": [
                    {"text": prompt},
                    {
                        "inline_data": {
                            "mime_type": mime_type,
                            "data": imagen_b64,
                        }
                    },
                ]
            }
        ]
    }

    for intento in range(3):
        respuesta = requests.post(URL_GEMINI, json=payload, timeout=120)

        if respuesta.status_code == 200:
            datos = respuesta.json()
            try:
                return datos["candidates"][0]["content"]["parts"][0]["text"]
            except (KeyError, IndexError):
                raise RuntimeError(f"Respuesta inesperada: {datos}")

        if respuesta.status_code == 503:
            print(f"  Gemini ocupado, reintentando ({intento + 1}/3)...")
            time.sleep(5)
            continue

        raise RuntimeError(f"Gemini error {respuesta.status_code}")

    raise RuntimeError("Gemini saturado tras 3 intentos.")


# ============ PROVEEDOR 2: GROQ ============

def transcribir_con_groq(imagen_b64, mime_type, prompt):
    headers = {
        "Authorization": f"Bearer {GROQ_API_KEY}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": MODELO_GROQ,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:{mime_type};base64,{imagen_b64}"
                        },
                    },
                ],
            }
        ],
        "temperature": 0.2,
    }

    respuesta = requests.post(URL_GROQ, headers=headers, json=payload, timeout=120)

    if respuesta.status_code != 200:
        raise RuntimeError(f"Groq error {respuesta.status_code}")

    datos = respuesta.json()
    try:
        return datos["choices"][0]["message"]["content"]
    except (KeyError, IndexError):
        raise RuntimeError(f"Respuesta inesperada de Groq: {datos}")


# ============ COORDINADOR ============

def transcribir_texto(imagen_b64, mime_type, prompt):
    """Intenta Gemini. Si falla, cae automáticamente a Groq."""
    try:
        print("  Proveedor principal: Gemini...")
        return transcribir_con_gemini(imagen_b64, mime_type, prompt)
    except Exception as e:
        print(f"  Gemini falló: {e}")
        print("  Cambiando al plan B: Groq...")
        try:
            return transcribir_con_groq(imagen_b64, mime_type, prompt)
        except Exception as e2:
            raise RuntimeError(
                f"Ambos proveedores fallaron.\n"
                f"  Gemini: {e}\n"
                f"  Groq: {e2}"
            )


# ============ GENERADORES DE SALIDA ============

def generar_docx(texto, ruta_salida, titulo="Documento generado por EscribIA"):
    doc = Document()
    doc.add_heading(titulo, level=1)
    for linea in texto.split("\n"):
        linea = linea.strip()
        if linea:
            doc.add_paragraph(linea)
    doc.save(ruta_salida)
    return ruta_salida


def parsear_slides(texto):
    """Convierte el texto estructurado en lista de diapositivas."""
    slides = []
    bloques = texto.split("---")

    for bloque in bloques:
        lineas = [l.strip() for l in bloque.strip().split("\n") if l.strip()]
        if not lineas:
            continue

        titulo = ""
        explicacion = ""
        puntos = []

        for linea in lineas:
            mayus = linea.upper()
            if mayus.startswith("TITULO:"):
                titulo = linea.split(":", 1)[1].strip()
            elif mayus.startswith("EXPLICACION:"):
                explicacion = linea.split(":", 1)[1].strip()
            elif linea.startswith("-"):
                puntos.append(linea[1:].strip())
            elif not titulo:
                titulo = linea
            elif not explicacion:
                explicacion = linea
            else:
                puntos.append(linea)

        if titulo or puntos:
            slides.append({
                "titulo": titulo or "Sin título",
                "explicacion": explicacion,
                "puntos": puntos,
            })

    return slides


def generar_pptx(texto, ruta_salida):
    """Genera una presentación con título, explicación y puntos por diapositiva."""
    slides = parsear_slides(texto)

    if not slides:
        slides = [{
            "titulo": "Apunte",
            "explicacion": "",
            "puntos": [texto[:200]],
        }]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Portada
    portada = prs.slides.add_slide(prs.slide_layouts[0])
    portada.shapes.title.text = "Apuntes digitalizados"
    portada.placeholders[1].text = (
        f"Generado por EscribIA el {datetime.now().strftime('%d/%m/%Y')}"
    )

    # Una diapositiva por sección
    for info in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = info["titulo"]

        if info["explicacion"]:
            caja_expl = slide.shapes.add_textbox(
                Inches(0.6), Inches(1.6), Inches(12.1), Inches(1.1)
            )
            tf = caja_expl.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = info["explicacion"]
            p.font.size = Pt(16)
            p.font.italic = True

        if info["puntos"]:
            top_puntos = Inches(3.0) if info["explicacion"] else Inches(2.0)
            caja_puntos = slide.shapes.add_textbox(
                Inches(0.6), top_puntos, Inches(12.1), Inches(4.2)
            )
            tf2 = caja_puntos.text_frame
            tf2.word_wrap = True

            for i, punto in enumerate(info["puntos"]):
                parrafo = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
                parrafo.text = f"•  {punto}"
                parrafo.font.size = Pt(20)

    prs.save(ruta_salida)
    return ruta_salida


# ============ PROGRAMA PRINCIPAL ============

def main():
    print("=" * 45)
    print("   EscribIA - Prototipo")
    print("=" * 45)
    print()

    print("¿Qué quieres generar a partir del apunte?")
    print("  1. Documento de Word")
    print("  2. Presentación de PowerPoint")
    print("  3. Resumen corto para estudiar")
    print()

    while True:
        opcion = input("Elige (1/2/3): ").strip()
        if opcion in ("1", "2", "3"):
            break
        print("Opción inválida. Intenta de nuevo.")

    if opcion == "1":
        prompt = PROMPT_WORD
        titulo_doc = "Documento generado por EscribIA"
    elif opcion == "2":
        prompt = PROMPT_PPT
        titulo_doc = "Presentación generada por EscribIA"
    else:
        prompt = PROMPT_RESUMEN
        titulo_doc = "Resumen de estudio generado por EscribIA"

    print()

    print("Buscando la última foto de tu cámara...")
    ruta = encontrar_ultima_foto()

    if not ruta:
        print("No encontré fotos en las carpetas típicas.")
        ruta = input("Escribe la ruta manualmente: ").strip().strip('"')
    else:
        print(f"Encontré: {ruta}")
        respuesta = input("¿Usar esta foto? (s/n): ").strip().lower()
        if respuesta != "s":
            ruta = input("Escribe la ruta manualmente: ").strip().strip('"')

    print()

    print("Cargando imagen...")
    try:
        imagen_b64, mime = cargar_imagen(ruta)
    except Exception as e:
        print(f"Error al cargar la imagen: {e}")
        return

    print("Procesando con IA (puede tardar unos segundos)...")
    try:
        texto = transcribir_texto(imagen_b64, mime, prompt)
    except Exception as e:
        print(f"Error: {e}")
        return

    print()
    print("--- CONTENIDO GENERADO ---")
    print(texto)
    print("--- FIN ---")
    print()

    sello = datetime.now().strftime("%Y%m%d_%H%M%S")

    try:
        if opcion == "1":
            archivo = f"escribia_word_{sello}.docx"
            generar_docx(texto, archivo, titulo_doc)
        elif opcion == "2":
            archivo = f"escribia_ppt_{sello}.pptx"
            generar_pptx(texto, archivo)
        else:
            archivo = f"escribia_resumen_{sello}.docx"
            generar_docx(texto, archivo, titulo_doc)

        print(f"Archivo generado: {archivo}")
        print(f"Ubicación: {os.path.abspath(archivo)}")
    except Exception as e:
        print(f"Error al generar el archivo: {e}")


if __name__ == "__main__":
    main()