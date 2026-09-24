"""
EscribIA - Versión web (Streamlit)
Genera documentos digitales a partir de fotos de apuntes manuscritos.
Usa Groq como proveedor principal (con Gemini como respaldo opcional).
"""

import base64
import io
from datetime import datetime

import requests
import streamlit as st
from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt


# ============ LEER CLAVES (tolerante a fallos) ============

def _leer_clave(nombre):
    """Intenta leer una clave desde Secrets o config.py sin crashear."""
    try:
        return st.secrets[nombre]
    except Exception:
        pass
    try:
        import config
        return getattr(config, nombre)
    except Exception:
        return None


GEMINI_API_KEY = _leer_clave("GEMINI_API_KEY")
GROQ_API_KEY = _leer_clave("GROQ_API_KEY")

# Verificar que al menos una clave exista
if not GEMINI_API_KEY and not GROQ_API_KEY:
    st.error(
        "⚠️ No hay ninguna API key configurada. "
        "Agrega GEMINI_API_KEY o GROQ_API_KEY en los Secrets de Streamlit."
    )
    st.stop()

# Si Gemini falta o está mal formada, la ignoramos
if GEMINI_API_KEY and not GEMINI_API_KEY.startswith("AIzaSy"):
    GEMINI_API_KEY = None  # formato inválido, se ignora


# ============ CONFIGURACIÓN ============

MODELO_GEMINI = "gemini-2.5-flash"
MODELO_GROQ = "meta-llama/llama-4-scout-17b-16e-instruct"

URL_GEMINI = (
    f"https://generativelanguage.googleapis.com/v1beta/models/"
    f"{MODELO_GEMINI}:generateContent"
)
URL_GROQ = "https://api.groq.com/openai/v1/chat/completions"


PROMPT_WORD = """Eres un asistente que digitaliza apuntes manuscritos en español.
Transcribe con fidelidad, corrige ortografía, organiza en secciones,
usa guiones para listas. Devuelve SOLO el texto, sin markdown."""


PROMPT_PPT = """Convierte este apunte manuscrito en diapositivas con explicaciones.
Formato EXACTO:
TITULO: <título corto>
EXPLICACION: <1-2 frases explicando el tema>
- punto 1
- punto 2
---
TITULO: <otro título>
EXPLICACION: <explicación>
- punto 1
---
Entre 4 y 8 diapositivas, 3-5 puntos cada una. Sin asteriscos."""


PROMPT_RESUMEN = """Crea un resumen corto (máx 300 palabras) del apunte.
Empieza con un párrafo introductorio, luego viñetas con conceptos clave,
y termina con 3 preguntas de repaso. Sin markdown."""


# ============ MOTOR DE IA ============

def _groq(imagen_b64, mime, prompt):
    """Llama a Groq (proveedor principal)."""
    if not GROQ_API_KEY:
        raise RuntimeError("No hay clave de Groq configurada.")
    headers = {"Authorization": f"Bearer {GROQ_API_KEY}"}
    payload = {
        "model": MODELO_GROQ,
        "messages": [{
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {"type": "image_url",
                 "image_url": {"url": f"data:{mime};base64,{imagen_b64}"}},
            ],
        }],
        "temperature": 0.2,
    }
    r = requests.post(URL_GROQ, headers=headers, json=payload, timeout=120)
    if r.status_code != 200:
        raise RuntimeError(f"Groq error {r.status_code}: {r.text[:200]}")
    return r.json()["choices"][0]["message"]["content"]


def _gemini(imagen_b64, mime, prompt):
    """Llama a Gemini (proveedor secundario)."""
    if not GEMINI_API_KEY:
        raise RuntimeError("Gemini no está disponible.")
    url = f"{URL_GEMINI}?key={GEMINI_API_KEY}"
    payload = {
        "contents": [{
            "parts": [
                {"text": prompt},
                {"inline_data": {"mime_type": mime, "data": imagen_b64}},
            ]
        }]
    }
    r = requests.post(url, json=payload, timeout=120)
    if r.status_code != 200:
        raise RuntimeError(f"Gemini error {r.status_code}")
    return r.json()["candidates"][0]["content"]["parts"][0]["text"]


def procesar(imagen_b64, mime, prompt):
    """Intenta con Groq primero. Si falla, prueba Gemini."""
    errores = []

    # Intento 1: Groq
    if GROQ_API_KEY:
        try:
            return _groq(imagen_b64, mime, prompt), "Groq"
        except Exception as e:
            errores.append(f"Groq: {e}")

    # Intento 2: Gemini (si tiene formato válido)
    if GEMINI_API_KEY:
        try:
            return _gemini(imagen_b64, mime, prompt), "Gemini"
        except Exception as e:
            errores.append(f"Gemini: {e}")

    raise RuntimeError(" | ".join(errores) or "Sin proveedores disponibles.")


# ============ GENERADORES DE ARCHIVOS ============

def docx_bytes(texto, titulo):
    doc = Document()
    doc.add_heading(titulo, level=1)
    for linea in texto.split("\n"):
        if linea.strip():
            doc.add_paragraph(linea.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def parsear_slides(texto):
    slides = []
    for bloque in texto.split("---"):
        lineas = [l.strip() for l in bloque.strip().split("\n") if l.strip()]
        if not lineas:
            continue
        titulo, explicacion, puntos = "", "", []
        for linea in lineas:
            mayus = linea.upper()
            if mayus.startswith("TITULO:"):
                titulo = linea.split(":", 1)[1].strip()
            elif mayus.startswith("EXPLICACION:"):
                explicacion = linea.split(":", 1)[1].strip()
            elif linea.startswith("-"):
                puntos.append(linea[1:].strip())
            elif not titulo:
                titulo = linea
            elif not explicacion:
                explicacion = linea
            else:
                puntos.append(linea)
        if titulo or puntos:
            slides.append((titulo or "Sin título", explicacion, puntos))
    return slides


def pptx_bytes(texto):
    slides = parsear_slides(texto)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    portada = prs.slides.add_slide(prs.slide_layouts[0])
    portada.shapes.title.text = "Apuntes digitalizados"
    portada.placeholders[1].text = (
        f"EscribIA - {datetime.now().strftime('%d/%m/%Y')}"
    )

    for titulo, explicacion, puntos in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = titulo

        if explicacion:
            caja = slide.shapes.add_textbox(
                Inches(0.6), Inches(1.6), Inches(12.1), Inches(1.1)
            )
            caja.text_frame.word_wrap = True
            p = caja.text_frame.paragraphs[0]
            p.text = explicacion
            p.font.size = Pt(16)
            p.font.italic = True

        if puntos:
            top = Inches(3.0) if explicacion else Inches(2.0)
            caja2 = slide.shapes.add_textbox(
                Inches(0.6), top, Inches(12.1), Inches(4.2)
            )
            caja2.text_frame.word_wrap = True
            for i, punto in enumerate(puntos):
                if i == 0:
                    par = caja2.text_frame.paragraphs[0]
                else:
                    par = caja2.text_frame.add_paragraph()
                par.text = f"•  {punto}"
                par.font.size = Pt(20)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ============ INTERFAZ ============

st.set_page_config(page_title="EscribIA", page_icon="📝", layout="centered")

st.title("📝 EscribIA")
st.caption("Apuntes escritos a mano, convertidos en documentos digitales con IA")

st.divider()

formato = st.radio(
    "¿Qué quieres generar?",
    ["📄 Documento de Word", "📊 Presentación de PowerPoint", "📚 Resumen de estudio"],
)

foto = st.file_uploader(
    "Sube o toma la foto del apunte",
    type=["jpg", "jpeg", "png", "webp"],
)

if foto:
    img = Image.open(foto)
    st.image(img, caption="Apunte cargado", use_container_width=True)

    if st.button("✨ Procesar con IA", type="primary", use_container_width=True):
        img.thumbnail((1600, 1600))
        buf = io.BytesIO()
        img.convert("RGB").save(buf, "JPEG", quality=85)
        b64 = base64.b64encode(buf.getvalue()).decode()
        mime = "image/jpeg"

        if "Word" in formato:
            prompt = PROMPT_WORD
        elif "PowerPoint" in formato:
            prompt = PROMPT_PPT
        else:
            prompt = PROMPT_RESUMEN

        with st.spinner("Procesando con inteligencia artificial..."):
            try:
                texto, proveedor = procesar(b64, mime, prompt)
            except Exception as e:
                st.error(f"Error: {e}")
                st.stop()

        st.success(f"✅ Procesado con {proveedor}")
        st.subheader("Contenido generado")
        st.text_area("Texto", texto, height=300, label_visibility="collapsed")

        if "Word" in formato:
            datos = docx_bytes(texto, "Documento generado por EscribIA")
            nombre = f"escribia_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            mime_dl = ("application/vnd.openxmlformats-officedocument."
                       "wordprocessingml.document")
        elif "PowerPoint" in formato:
            datos = pptx_bytes(texto)
            nombre = f"escribia_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
            mime_dl = ("application/vnd.openxmlformats-officedocument."
                       "presentationml.presentation")
        else:
            datos = docx_bytes(texto, "Resumen de estudio - EscribIA")
            nombre = f"escribia_resumen_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            mime_dl = ("application/vnd.openxmlformats-officedocument."
                       "wordprocessingml.document")

        st.download_button(
            "⬇️ Descargar archivo",
            data=datos,
            file_name=nombre,
            mime=mime_dl,
            type="primary",
            use_container_width=True,
        )
